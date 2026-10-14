"""Executable self-check without reading user config or touching user PPT files."""
import json
import tempfile
import traceback
from pathlib import Path


def smoke_test(report_path):
    report={'ok':False}
    root=None
    try:
        import tkinter as tk
        from PIL import Image
        from pptx import Presentation
        import pythoncom
        import win32com.client
        from result_grid import ResultGrid
        from naver_import import PREFIX,parse_naver_text,group_naver,create_guide
        from web_import import create_guide as create_sise_guide
        root=tk.Tk();root.withdraw()
        import importlib.util
        import sys
        if getattr(sys, 'frozen', False):
            for module in ('ocr_engine','rapidocr_onnxruntime','onnxruntime','cv2','numpy','winocr'):
                assert importlib.util.find_spec(module) is None, module + ' must not be bundled'
        grid=ResultGrid(root);grid.set_text('단지 24\n84) 12.5/6.5')
        assert not grid.issues()
        assert grid.get_text().endswith('84) 12.5/6.5')
        Image.new('RGB',(20,20))
        data={'source':'https://new.land.naver.com/','complexId':'1','name':'검증','households':'780','approvalDate':'2024-01-30',
              'complete':True,'types':[{'key':'116A','supplyPyeong':'35','exclusiveArea':'84.15','sale':'12.5','lease':'6.5'}]}
        assert group_naver(parse_naver_text(PREFIX+json.dumps(data)))[0]['sale']=='12.5'
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'test.pptx';prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(p)
            assert len(Presentation(p).slides)==1
            assert 'javascript:' in create_guide(d).read_text('utf-8')
            create_sise_guide(d)
        report.update(ok=True,checks=['Tk/result-grid','PPTX','Pillow','COM imports','Naver parser','both bookmarklets'])
    except Exception:
        report['error']=traceback.format_exc()
    finally:
        if root is not None:root.destroy()
        Path(report_path).write_text(json.dumps(report,ensure_ascii=False,indent=2),'utf-8')
    return 0 if report['ok'] else 1
