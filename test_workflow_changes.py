import os
import tempfile
import unittest
from unittest import mock
from pptx import Presentation
from pptx.util import Cm
from config_manager import Config
from record_writer import result_to_entries, ZoneRecorder
import pptx_writer


class WorkflowChangesTest(unittest.TestCase):
    def test_record_boundaries(self):
        areas = [54,55,60,65,66,72,79,80,85,90,91]
        text = "단지 99\n" + "\n".join(f"{a}) 1/2" for a in areas)
        self.assertEqual([a for a, _ in result_to_entries(text)], ["24"]*3 + ["34"]*3)
        with tempfile.TemporaryDirectory() as d:
            r = ZoneRecorder(d)
            r.add("가 99\n60) 1/2")
            r.add("나 99\n72) 3/2")
            r.replace_last("나 99\n85) 4/2")
            self.assertEqual(sorted(os.listdir(d)), ["미지정_24평.txt", "미지정_34평.txt"])
            from pathlib import Path
            self.assertIn("가", Path(r.file_for("24")).read_text(encoding="utf-8-sig"))

    def test_titles_alias_and_missing_slide(self):
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "test.pptx")
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_textbox(0, 0, Cm(10), Cm(1)).text = "선사/천호 생활권 "
            sid = s.slide_id
            prs.slides.add_slide(prs.slide_layouts[6])
            prs.save(path)
            with mock.patch.object(pptx_writer, "_find_open_presentation", return_value=None):
                rows = pptx_writer.read_zone_titles(path)
                self.assertEqual(rows[0]["name"], "선사/천호 생활권")
                self.assertEqual(rows[0]["slide_id"], sid)
                cfg = Config(pptx_path=path, zone_slide_id=sid)
                self.assertEqual(pptx_writer.add_text_to_pptx("가 99", cfg, zone="수정 이름")[0], 1)
                cfg.zone_slide_id = 999999
                with self.assertRaises(ValueError):
                    pptx_writer.add_text_to_pptx("가 99", cfg, zone="수정 이름")
            self.assertEqual(len(Presentation(path).slides), 2)

    def test_zone_editor_and_raw_area_recording(self):
        import tkinter as tk
        import gui_app
        from zone_dialog import ZoneDialog
        root = tk.Tk(); root.withdraw()
        try:
            with tempfile.TemporaryDirectory() as d, mock.patch.object(gui_app, "save_config"):
                cfg = Config(pptx_path=os.path.join(d,"missing.pptx"), zone="가 생활권", result_in_pyeong=True,
                             auto_copy=False, use_name_dict=False)
                gui = gui_app.SisePickerGUI(root, config=cfg, ask_zone_on_start=False)
                gui.result_grid.set_text("단지 99\n55) 1/2\n65) 3/2\n72) 4/2\n80) 5/2\n90) 6/2")
                gui._record(gui.result_text(), True)
                self.assertEqual(sorted(os.listdir(gui.recorder.directory)), ["가 생활권_24평.txt", "가 생활권_34평.txt"])
                saved=[]
                dialog = ZoneDialog(root, [dict(name="가", slide_id=3, page=1)], lambda: [], lambda rows, name: saved.append((rows,name)))
                dialog.tree.selection_set("0"); dialog.name.set("변경"); dialog.modify()
                self.assertEqual(dialog.rows[0]["slide_id"], 3)
                dialog.name.set("새 생활권"); dialog.add()
                self.assertEqual(len(dialog.rows), 2)
                dialog.remove(); self.assertEqual(len(dialog.rows), 1)
                dialog.tree.selection_set("0"); dialog.save()
                self.assertEqual(saved[0][1], "변경")
                gui.on_close()
        finally:
            try: root.destroy()
            except tk.TclError: pass

if __name__ == "__main__":
    unittest.main()
