import json
import tempfile
import tkinter as tk
import unittest
from pathlib import Path
from unittest import mock
from config_manager import Config
import gui_app
from web_import import PREFIX
from pptx import Presentation

class WebGuiTest(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=tk.Tk(); self.root.withdraw()
        self.save=mock.patch.object(gui_app,'save_config');self.save.start()
        self.errors=mock.patch.object(gui_app.messagebox,'showerror',side_effect=AssertionError('Unexpected error dialog'));self.errors.start()
        self.gui=gui_app.SisePickerGUI(self.root,Config(pptx_path=str(Path(self.tmp.name)/'test.pptx'),zone='테스트 생활권'),ask_zone_on_start=False)
    def tearDown(self):
        self.gui.on_close();self.save.stop();self.errors.stop();self.tmp.cleanup()
    def test_tabs_and_clipboard_routes_without_ocr(self):
        g=self.gui
        self.assertEqual([g.notebook.tab(t,'text') for t in g.notebook.tabs()],['네이버 부동산','시세통'])
        self.assertTrue(g._nav_mode())
        self.assertFalse(hasattr(g,'process_image'))
        with mock.patch.object(g,'_read_naver_clipboard',return_value=True) as nav, mock.patch.object(g,'_read_web_clipboard',return_value=True) as web:
            g.read_clipboard();nav.assert_called_once();web.assert_not_called()
            g.notebook.select(g.web_tab);g.read_clipboard();web.assert_called_once()
        with mock.patch.object(self.root,'clipboard_get',side_effect=tk.TclError('image only')):
            g.read_clipboard()
        self.assertIn('시세통 전용',g.var_status.get())
    def test_import_edit_ppt_and_repeat_watch_preserves_edits(self):
        g=self.gui;g.notebook.select(g.web_tab)
        raw=PREFIX+json.dumps(dict(source='https://sisetong.kr/',name='예시',year='19',areas=['60','85'],rows=[['6.0','9.0'],['4.0','5.0']]))
        with mock.patch.object(self.root,'clipboard_get',return_value=raw):
            g.read_clipboard()
            self.assertFalse(Path(g.config.pptx_path).exists())
            g.result_grid.rows[0][0][1].set('6.5')
            g._read_web_clipboard(watch=True)
            self.assertIn('60) 6.5/4.0',g.result_text())
        template=Presentation();template.slides.add_slide(template.slide_layouts[6]);template.save(g.config.pptx_path)
        self.assertTrue(g._ready_for_output(), (g.result_text(), g.result_grid.issues()))
        with mock.patch('pptx_writer._find_open_presentation',return_value=None):
            g.add_to_ppt()
        prs=Presentation(g.config.pptx_path)
        self.assertTrue(any('60) 6.5/4.0' in sh.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame))
        self.assertTrue(Path(g.recorder.file_for('24')).exists())
    def test_invalid_price_blocks_output(self):
        g=self.gui;g.result_grid.set_text('예시 19\n60) ?/4.0')
        self.assertFalse(g._ready_for_output())

if __name__=='__main__':unittest.main()
